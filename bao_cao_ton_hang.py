import sys
import os
import io
import re
import csv
from datetime import datetime

from PyQt5 import QtWidgets, uic
from PyQt5.QtWidgets import (
    QMessageBox, QTableWidgetItem, QHeaderView, QComboBox,
    QFileDialog, QPushButton, QLineEdit, QLabel, QTableWidget
)
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QFont, QColor, QBrush
import pymysql

# ===================== KẾT NỐI MYSQL =====================
DB_CONFIG = {
    "host": "localhost",
    "port": 3306,
    "user": "root",
    "password": "123456789",
    "database": "Cafe",
    "charset": "utf8mb4",
    "cursorclass": pymysql.cursors.DictCursor,
}

# ===================== CỠ CHỮ (pt) =====================
CO_CHU_DAU_COT = 12     # tiêu đề cột
CO_CHU_O = 10           # các dòng dữ liệu (nhỏ hơn tiêu đề 2 cỡ)
CO_CHU_NUT = 12
CAO_DONG_BANG = 52

UI_FILE = "bao_cao_ton_hang.ui"

# Mức tồn tối thiểu mặc định (theo đơn vị tính) khi tự thêm cột lần đầu
MUC_TON_MAC_DINH = {"kg": 5, "lít": 10, "gói": 5, "lon": 10}
MUC_TON_KHAC = 5

TT_CAN_NHAP = "Cần nhập"
TT_SAN_CO = "Sẵn có"


def doc_ui_tuong_thich(duong_dan):
    """Đọc file .ui và đổi kiểu enum của Qt6 (QFrame::Shape::StyledPanel,
    Qt::Orientation::Vertical...) về kiểu PyQt5 hiểu được (QFrame::StyledPanel...).
    Nhờ vậy lưu file .ui bằng Qt Designer bản nào cũng chạy được."""
    with open(duong_dan, encoding="utf-8") as f:
        noi_dung = f.read()

    def sua(m):
        return m.group(1) + re.sub(r"(\w+)::\w+::(\w+)", r"\1::\2", m.group(2)) + m.group(3)

    noi_dung = re.sub(r"(<(?:enum|set)>)(.*?)(</(?:enum|set)>)", sua, noi_dung, flags=re.S)
    # Cỡ chữ -1 (không hợp lệ) làm Qt báo "Point size <= 0" -> bỏ đi để dùng cỡ mặc định
    noi_dung = re.sub(r"\s*<pointsize>-1</pointsize>", "", noi_dung)
    return io.BytesIO(noi_dung.encode("utf-8"))


def them_style(widget, qss):
    cu = widget.styleSheet().strip()
    if cu and "{" not in cu:
        cu = "* { " + cu + " }"
    widget.setStyleSheet(cu + "\n" + qss)


def dinh_dang_so(so):
    """5.000 -> 5, 2.500 -> 2.5 (bỏ số 0 thừa của DECIMAL(12,3))."""
    return f"{float(so or 0):.3f}".rstrip("0").rstrip(".") or "0"


class BaoCaoTonHangWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        base_dir = os.path.dirname(os.path.abspath(__file__))
        os.chdir(base_dir)
        uic.loadUi(doc_ui_tuong_thich(os.path.join(base_dir, UI_FILE)), self)

        self.tat_ca = []        # toàn bộ nguyên liệu đọc từ DB
        self.dang_hien = []     # các dòng đang hiển thị (sau khi lọc) -> dùng để xuất Excel

        self.lay_widget()
        self.cau_hinh_giao_dien()
        self.noi_su_kien()

        self.dam_bao_cot_muc_ton()
        self.tai_du_lieu()

    # ------------------------------------------------------------------
    def lay_widget(self):
        def w(loai, ten):
            obj = self.findChild(loai, ten)
            if obj is None:
                raise RuntimeError(f"Không tìm thấy widget '{ten}' trong {UI_FILE}.")
            return obj

        self.cbo_tt = w(QComboBox, "comboBox")
        self.txt_tim = w(QLineEdit, "txtTimNL")
        self.btn_loc = w(QPushButton, "btnLocDuLieu")
        self.btn_excel = w(QPushButton, "btnXuatExcel")
        self.lbl_tong = w(QLabel, "lbl_tongNL")
        self.lbl_can_nhap = w(QLabel, "lbl_CanNhap")
        self.lbl_san_co = w(QLabel, "lbl_SanCo")
        self.bang = w(QTableWidget, "tableTonKho")

    def cau_hinh_giao_dien(self):
        # Ô tìm kiếm: xóa chữ gõ sẵn, chỉ để gợi ý mờ
        self.txt_tim.clear()
        self.txt_tim.setPlaceholderText("Nhập mã hoặc tên nguyên liệu...")

        # Nút
        for b in (self.btn_loc, self.btn_excel):
            b.setMaximumWidth(16777215)
            b.setMinimumSize(150, 44)
            f = b.font()
            f.setPointSize(CO_CHU_NUT)
            f.setBold(True)
            b.setFont(f)

        # Bảng
        self.bang.setColumnCount(7)
        self.bang.setHorizontalHeaderLabels([
            "STT", "Mã nguyên liệu", "Tên nguyên liệu", "Đơn vị tính",
            "Số lượng tồn", "Mức tồn tối thiểu", "Cảnh báo"
        ])
        self.bang.verticalHeader().setVisible(False)
        self.bang.setAlternatingRowColors(True)
        self.bang.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.bang.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.bang.setSelectionMode(QtWidgets.QAbstractItemView.SingleSelection)
        self.bang.setWordWrap(False)

        them_style(self.bang, f"""
            QTableWidget {{ font-size: {CO_CHU_O}pt; }}
            QHeaderView::section {{ font-size: {CO_CHU_DAU_COT}pt; }}
            QTableWidget::item {{ padding: 5px; font-size: {CO_CHU_O}pt; }}
        """)
        f_o = QFont(self.bang.font())
        f_o.setPointSize(CO_CHU_O)
        self.bang.setFont(f_o)
        f_dau = QFont(f_o)
        f_dau.setPointSize(CO_CHU_DAU_COT)
        f_dau.setBold(True)
        self.bang.horizontalHeader().setFont(f_dau)

        h = self.bang.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(2, QHeaderView.Stretch)
        h.setSectionResizeMode(3, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(4, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(5, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(6, QHeaderView.Stretch)

    def noi_su_kien(self):
        self.btn_loc.clicked.connect(self.bam_loc)
        self.btn_excel.clicked.connect(self.xuat_excel)
        self.txt_tim.returnPressed.connect(self.bam_loc)
        self.cbo_tt.currentIndexChanged.connect(self.ap_dung_loc)

        # Menu bên trái
        def noi(ten, ham):
            b = self.findChild(QPushButton, ten)
            if b is not None:
                b.clicked.connect(ham)
        noi("btnDonHang", self.mo_lap_don_hang)
        noi("btnNguyenLieu", self.mo_nhap_nguyen_lieu)
        noi("btnTimKiemDonHang", self.mo_tim_kiem)

    # ------------------------------------------------------------------
    # DATABASE
    def ket_noi(self):
        return pymysql.connect(**DB_CONFIG)

    def dam_bao_cot_muc_ton(self):
        """Bảng nguyen_lieu chưa có cột 'Mức tồn tối thiểu' -> tự thêm lần đầu."""
        try:
            conn = self.ket_noi()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi kết nối", f"Không thể kết nối CSDL:\n{e}")
            return
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """SELECT COUNT(*) AS n FROM INFORMATION_SCHEMA.COLUMNS
                       WHERE TABLE_SCHEMA = DATABASE()
                         AND TABLE_NAME = 'nguyen_lieu'
                         AND COLUMN_NAME = 'MucTonToiThieu'"""
                )
                if cur.fetchone()["n"] == 0:
                    cur.execute(
                        "ALTER TABLE nguyen_lieu "
                        "ADD COLUMN MucTonToiThieu DECIMAL(12,3) NOT NULL DEFAULT 0"
                    )
                    for dv, muc in MUC_TON_MAC_DINH.items():
                        cur.execute(
                            "UPDATE nguyen_lieu SET MucTonToiThieu = %s WHERE DonViTinh = %s",
                            (muc, dv)
                        )
                    cur.execute(
                        "UPDATE nguyen_lieu SET MucTonToiThieu = %s WHERE MucTonToiThieu = 0",
                        (MUC_TON_KHAC,)
                    )
            conn.commit()
        except Exception as e:
            conn.rollback()
            QMessageBox.critical(self, "Lỗi CSDL", f"Không thêm được cột mức tồn tối thiểu:\n{e}")
        finally:
            conn.close()

    def tai_du_lieu(self):
        try:
            conn = self.ket_noi()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi kết nối", f"Không thể kết nối CSDL:\n{e}")
            return
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """SELECT MaNL, TenNL, DonViTinh, SoLuongTon, MucTonToiThieu
                       FROM nguyen_lieu ORDER BY MaNL"""
                )
                self.tat_ca = cur.fetchall()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi Database", f"Lỗi khi đọc tồn kho:\n{e}")
            self.tat_ca = []
        finally:
            conn.close()

        for r in self.tat_ca:
            r["TrangThai"] = self.tinh_trang_thai(r)

        self.cap_nhat_the_thong_ke()
        self.ap_dung_loc()

    @staticmethod
    def tinh_trang_thai(r):
        ton = float(r["SoLuongTon"] or 0)
        toi_thieu = float(r["MucTonToiThieu"] or 0)
        return TT_CAN_NHAP if ton < toi_thieu else TT_SAN_CO

    # ------------------------------------------------------------------
    # LỌC + HIỂN THỊ
    def cap_nhat_the_thong_ke(self):
        can = sum(1 for r in self.tat_ca if r["TrangThai"] == TT_CAN_NHAP)
        self.lbl_tong.setText(str(len(self.tat_ca)))
        self.lbl_can_nhap.setText(str(can))
        self.lbl_san_co.setText(str(len(self.tat_ca) - can))

    def bam_loc(self):
        """Bấm 'Lọc dữ liệu' -> đọc lại DB cho chắc có số liệu mới nhất rồi lọc."""
        self.tai_du_lieu()

    def ap_dung_loc(self):
        tu_khoa = self.txt_tim.text().strip().lower()
        chon = self.cbo_tt.currentText()
        ds = []
        for r in self.tat_ca:
            if tu_khoa and tu_khoa not in str(r["MaNL"]).lower() \
                    and tu_khoa not in str(r["TenNL"]).lower():
                continue
            if chon == TT_SAN_CO and r["TrangThai"] != TT_SAN_CO:
                continue
            if chon == TT_CAN_NHAP and r["TrangThai"] != TT_CAN_NHAP:
                continue
            ds.append(r)
        self.dang_hien = ds
        self.do_du_lieu(ds)

    def do_du_lieu(self, rows):
        f_o = QFont(self.bang.font())
        f_o.setPointSize(CO_CHU_O)

        self.bang.setRowCount(0)
        for i, d in enumerate(rows):
            self.bang.insertRow(i)
            self.bang.setRowHeight(i, CAO_DONG_BANG)

            gia_tri = [
                str(i + 1), str(d["MaNL"]), d["TenNL"] or "", d["DonViTinh"] or "",
                dinh_dang_so(d["SoLuongTon"]), dinh_dang_so(d["MucTonToiThieu"]),
                d["TrangThai"],
            ]
            for c, text in enumerate(gia_tri):
                item = QTableWidgetItem(text)
                item.setFont(f_o)
                if c == 2:
                    item.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
                elif c in (4, 5):
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                else:
                    item.setTextAlignment(Qt.AlignCenter)

                if c == 6:
                    if d["TrangThai"] == TT_CAN_NHAP:
                        item.setForeground(QBrush(QColor("#A33C35")))
                        item.setBackground(QBrush(QColor("#F9E3E0")))
                    else:
                        item.setForeground(QBrush(QColor("#2E7D32")))
                        item.setBackground(QBrush(QColor("#E6F2E6")))
                    bold = QFont(f_o)
                    bold.setBold(True)
                    item.setFont(bold)
                self.bang.setItem(i, c, item)

    # ------------------------------------------------------------------
    # XUẤT EXCEL
    def xuat_excel(self):
        if not self.dang_hien:
            QMessageBox.information(self, "Xuất Excel", "Không có dữ liệu để xuất.")
            return

        mac_dinh = f"bao_cao_ton_kho_{datetime.now():%Y%m%d_%H%M}.xlsx"
        duong_dan, _ = QFileDialog.getSaveFileName(
            self, "Lưu báo cáo tồn kho", mac_dinh, "Excel (*.xlsx)"
        )
        if not duong_dan:
            return
        if not duong_dan.lower().endswith(".xlsx"):
            duong_dan += ".xlsx"

        tieu_de = ["STT", "Mã nguyên liệu", "Tên nguyên liệu", "Đơn vị tính",
                   "Số lượng tồn", "Mức tồn tối thiểu", "Cảnh báo"]
        dong = []
        for i, d in enumerate(self.dang_hien, start=1):
            dong.append([
                i, d["MaNL"], d["TenNL"], d["DonViTinh"],
                float(d["SoLuongTon"] or 0), float(d["MucTonToiThieu"] or 0),
                d["TrangThai"],
            ])

        try:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
            except ImportError:
                # Máy chưa cài openpyxl -> xuất CSV (Excel vẫn mở được)
                duong_csv = duong_dan[:-5] + ".csv"
                with open(duong_csv, "w", newline="", encoding="utf-8-sig") as f:
                    w = csv.writer(f)
                    w.writerow(tieu_de)
                    w.writerows(dong)
                QMessageBox.information(
                    self, "Đã xuất file CSV",
                    f"Máy chưa cài openpyxl nên đã xuất file CSV (mở bằng Excel được):\n{duong_csv}\n\n"
                    "Muốn xuất .xlsx: mở CMD gõ  pip install openpyxl"
                )
                return

            wb = Workbook()
            ws = wb.active
            ws.title = "Tồn kho"
            ws.append(tieu_de)
            for d in dong:
                ws.append(d)

            for c in ws[1]:
                c.font = Font(bold=True, color="493020")
                c.fill = PatternFill("solid", fgColor="EDE1D4")
                c.alignment = Alignment(horizontal="center", vertical="center")
            for hang in ws.iter_rows(min_row=2):
                if hang[6].value == TT_CAN_NHAP:
                    hang[6].font = Font(bold=True, color="A33C35")
                else:
                    hang[6].font = Font(bold=True, color="2E7D32")
            for cot, rong in zip("ABCDEFG", (8, 18, 30, 14, 16, 20, 14)):
                ws.column_dimensions[cot].width = rong
            wb.save(duong_dan)
            QMessageBox.information(self, "Xuất Excel", f"Đã xuất báo cáo:\n{duong_dan}")
        except PermissionError:
            QMessageBox.critical(
                self, "Không lưu được",
                "File đang được mở bằng Excel. Hãy đóng file đó rồi thử lại."
            )
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất file", str(e))

    # ------------------------------------------------------------------
    # CHUYỂN MÀN HÌNH (menu bên trái)
    def _chuyen(self, lop):
        self.cua_so = lop()
        self.cua_so.setGeometry(self.geometry())
        self.cua_so.show()
        self.close()

    def mo_lap_don_hang(self):
        from lap_don_hang import LapDonHangWindow
        self._chuyen(LapDonHangWindow)

    def mo_nhap_nguyen_lieu(self):
        from main import NhapNguyenLieuWindow
        self._chuyen(NhapNguyenLieuWindow)

    def mo_tim_kiem(self):
        from tim_kiem_don_hang import TimKiemDonHangWindow
        self._chuyen(TimKiemDonHangWindow)


if __name__ == "__main__":
    app = QtWidgets.QApplication(sys.argv)
    window = BaoCaoTonHangWindow()
    window.show()
    sys.exit(app.exec_())